"""Safely extract PDF, UTF-8 TXT, and DOCX text with genuine source metadata.

Files stay in memory. DOCX archive expansion and extracted text are bounded.
Scanned PDFs require OCR, which is intentionally outside this capstone.
"""

import io
import re
import zipfile
from pathlib import PurePosixPath

from docx import Document as WordDocument
from langchain_core.documents import Document
from pypdf import PdfReader

from backend.core.exceptions import AppError

SUPPORTED_TYPES = {".txt", ".pdf", ".docx"}


def safe_filename(name: str) -> str:
    """Remove both Windows/POSIX path components and unsafe display characters."""
    basename = PurePosixPath(name.replace("\\", "/")).name
    cleaned = re.sub(r"[^a-zA-Z0-9._ -]", "_", basename).strip(" .")
    if not cleaned or len(cleaned) > 150:
        raise AppError(422, "invalid_filename", "Use a filename of 1–150 safe characters.")
    return cleaned


def load_document(name: str, data: bytes, document_id: str, max_chars: int) -> list[Document]:
    """Extract bounded text or raise a safe error for corrupt/empty input.

    PDF pages are one-based. TXT and DOCX have no reliable pagination, so no
    page number is manufactured. DOCX tables are included in extraction.
    """
    extension = PurePosixPath(name).suffix.lower()
    if extension not in SUPPORTED_TYPES:
        raise AppError(415, "unsupported_document", "Supported file types are PDF, TXT, and DOCX.")
    if not data:
        raise AppError(422, "empty_document", "The uploaded file is empty.")
    parts: list[tuple[str, int | None]] = []
    try:
        if extension == ".txt":
            text = data.decode("utf-8-sig")
            if "\x00" in text:
                raise ValueError("Binary text")
            parts = [(text, None)]
        elif extension == ".pdf":
            if not data.startswith(b"%PDF-"):
                raise ValueError("Missing PDF signature")
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise AppError(422, "encrypted_pdf", "Upload an unencrypted PDF.")
            total = 0
            for page_number, page in enumerate(reader.pages, start=1):
                text = page.extract_text() or ""
                total += len(text)
                if total > max_chars:
                    raise AppError(413, "document_too_long", "Extracted text exceeds the limit.")
                parts.append((text, page_number))
        else:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if "word/document.xml" not in archive.namelist():
                    raise ValueError("Not a DOCX package")
                if sum(item.file_size for item in archive.infolist()) > 30 * 1024 * 1024:
                    raise AppError(413, "archive_too_large", "DOCX expanded content is too large.")
            document = WordDocument(io.BytesIO(data))
            text_parts = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                text_parts.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
            parts = [("\n".join(text_parts), None)]
    except AppError:
        raise
    except Exception:
        raise AppError(
            422,
            "unreadable_document",
            "Cannot read this file. Use a valid PDF, DOCX, or UTF-8 TXT document.",
        ) from None
    if sum(len(text) for text, _ in parts) > max_chars:
        raise AppError(413, "document_too_long", "Extracted text exceeds the configured limit.")
    result = []
    for text, page in parts:
        if text.strip():
            result.append(
                Document(
                    page_content=text.strip(),
                    metadata={
                        "file_name": name,
                        "source": name,
                        "page": page,
                        "document_id": document_id,
                    },
                )
            )
    if not result:
        raise AppError(
            422,
            "empty_document",
            "No readable text was found. Scanned PDFs need OCR before uploading.",
        )
    return result
