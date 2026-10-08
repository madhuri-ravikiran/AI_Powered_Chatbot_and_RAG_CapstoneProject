"""Check genuine TXT/PDF/DOCX extraction, provenance, and invalid files."""

import io
import zipfile

import pytest
from docx import Document as WordDocument
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from backend.ai.document_loader import load_document, safe_filename
from backend.core.exceptions import AppError


def text_pdf() -> bytes:
    """Create a one-page real PDF fixture with searchable text, without extra dependencies."""
    writer = PdfWriter()
    page = writer.add_blank_page(width=600, height=800)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 50 700 Td (Annual leave is 24 days.) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def test_txt_metadata():
    """UTF-8 text keeps a filename and never invents page numbers."""
    docs = load_document("policy.txt", b"Annual leave is 24 days.", "abc", 10000)
    assert docs[0].page_content == "Annual leave is 24 days."
    assert docs[0].metadata["page"] is None
    assert docs[0].metadata["source"] == "policy.txt"


def test_pdf_extraction():
    """Read a real generated PDF and preserve one-based page metadata."""
    docs = load_document("policy.pdf", text_pdf(), "abc", 10000)
    assert "24 days" in docs[0].page_content
    assert docs[0].metadata["page"] == 1


def test_docx_paragraphs_and_tables():
    """Extract DOCX tables as well as paragraphs without guessing pagination."""
    document = WordDocument()
    document.add_paragraph("Annual leave policy")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "24 days"
    output = io.BytesIO()
    document.save(output)
    docs = load_document("policy.docx", output.getvalue(), "abc", 10000)
    assert "24 days" in docs[0].page_content
    assert docs[0].metadata["page"] is None


@pytest.mark.parametrize(
    ("name", "data", "code"),
    [
        ("empty.txt", b"", "empty_document"),
        ("blank.txt", b" \n ", "empty_document"),
        ("bad.txt", b"\xff", "unreadable_document"),
        ("binary.txt", b"x\x00y", "unreadable_document"),
        ("bad.pdf", b"hello", "unreadable_document"),
        ("bad.docx", b"not zip", "unreadable_document"),
        ("code.exe", b"hello", "unsupported_document"),
    ],
)
def test_invalid_documents(name, data, code):
    """Reject invalid content with safe stable errors."""
    with pytest.raises(AppError) as error:
        load_document(name, data, "abc", 10000)
    assert error.value.code == code


def test_filename_sanitization_and_text_limit():
    """Remove traversal components and reject excessive extracted text."""
    assert safe_filename("../private/policy.txt") == "policy.txt"
    assert safe_filename(r"C:\private\policy.txt") == "policy.txt"
    with pytest.raises(AppError):
        load_document("big.txt", b"x" * 101, "abc", 100)


def test_docx_archive_limit():
    """Reject oversized expanded archives before parsing XML."""
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"x" * (31 * 1024 * 1024))
    with pytest.raises(AppError) as error:
        load_document("bomb.docx", output.getvalue(), "abc", 10000)
    assert error.value.code == "archive_too_large"
